"""Download the official UCI workbook and check stage-one feasibility.

Run from any directory: python data-analysis-project/docs/check_data.py
Requires openpyxl==3.1.5. No cleaning or final analysis is performed.
"""

from collections import Counter
from datetime import datetime, timezone
from hashlib import sha256
from importlib.metadata import version
import json
from pathlib import Path
import shutil
import sys
from urllib.request import urlopen
from zipfile import ZipFile

from openpyxl import load_workbook


PROJECT = Path(__file__).resolve().parents[1]
SOURCE_URL = "https://archive.ics.uci.edu/dataset/352/online-retail"
DOWNLOAD_URL = "https://archive.ics.uci.edu/static/public/352/online%2Bretail.zip"
EXPECTED_COLUMNS = [
    "InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate",
    "UnitPrice", "CustomerID", "Country",
]


def main():
    raw = PROJECT / "data" / "raw"
    raw.mkdir(parents=True, exist_ok=True)
    archive = raw / "online_retail.zip"
    workbook = raw / "Online Retail.xlsx"
    if not workbook.exists():
        if not archive.exists():
            print("Downloading official UCI archive...", flush=True)
            temporary = archive.with_suffix(".zip.part")
            with urlopen(DOWNLOAD_URL, timeout=90) as response:
                with temporary.open("wb") as output:
                    shutil.copyfileobj(response, output)
            with ZipFile(temporary) as zf:
                if zf.testzip() is not None:
                    raise ValueError("Archive CRC check failed")
                zf.getinfo("Online Retail.xlsx")
            temporary.replace(archive)
        with ZipFile(archive) as zf:
            temporary = workbook.with_suffix(".xlsx.part")
            with zf.open("Online Retail.xlsx") as source:
                with temporary.open("wb") as output:
                    shutil.copyfileobj(source, output)
            temporary.replace(workbook)

    digest = sha256()
    with workbook.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    book = load_workbook(workbook, read_only=True, data_only=True)
    sheet = book.active
    rows = sheet.iter_rows(values_only=True)
    columns = list(next(rows))
    if columns != EXPECTED_COLUMNS:
        raise ValueError(f"Unexpected columns: {columns}")

    missing = Counter({name: 0 for name in columns})
    seen = set()
    checks = Counter()
    scope = Counter()
    months = set()
    first_date = last_date = None
    for row in rows:
        checks["rows"] += 1
        for name, value in zip(columns, row):
            if value is None:
                missing[name] += 1
        if row in seen:
            checks["duplicate_rows_after_first"] += 1
        seen.add(row)
        invoice, stock, description, quantity, date, price, customer, country = row
        cancellation = str(invoice).upper().startswith("C")
        checks["cancellation_rows"] += int(cancellation)
        checks["nonpositive_quantity_rows"] += int(quantity is not None and quantity <= 0)
        checks["nonpositive_unit_price_rows"] += int(price is not None and price <= 0)
        if not isinstance(date, datetime):
            checks["unreadable_date_rows"] += 1
            continue
        first_date = date if first_date is None else min(first_date, date)
        last_date = date if last_date is None else max(last_date, date)
        if country == "United Kingdom" and datetime(2011, 1, 1) <= date < datetime(2011, 12, 1):
            scope["rows_before_cleaning"] += 1
            scope["cancellation_rows"] += int(cancellation)
            months.add(date.strftime("%Y-%m"))

    report = {
        "checked_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "stage": "01_proposal: raw-data feasibility check; no cleaning applied",
        "source_url": SOURCE_URL,
        "download_url": DOWNLOAD_URL,
        "citation": "Chen, D. (2015). Online Retail. UCI Machine Learning Repository. https://doi.org/10.24432/C5BW33",
        "license": "CC BY 4.0",
        "workbook": str(workbook.relative_to(PROJECT)),
        "bytes": workbook.stat().st_size,
        "sha256": digest.hexdigest(),
        "sheet": sheet.title,
        "shape": [checks["rows"], len(columns)],
        "columns": columns,
        "date_range": [first_date.isoformat(), last_date.isoformat()],
        "missing_cells": dict(missing),
        "raw_quality_counts": dict(checks),
        "planned_scope": {
            "country": "United Kingdom",
            "start_inclusive": "2011-01-01",
            "end_exclusive": "2011-12-01",
            **dict(scope),
            "months_present": sorted(months),
        },
        "environment": {
            "python": sys.version.split()[0],
            "openpyxl": version("openpyxl"),
        },
    }
    book.close()
    destination = PROJECT / "docs" / "data_check.json"
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
